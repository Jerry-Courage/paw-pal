import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from django.test import SimpleTestCase, TestCase, override_settings
from library.source_understanding import build_understanding, grounding_bundle, validate_semantics, understand_with_ai, persist_understanding
from library.text_extractor import extract_text_from_bytes
from .material_grounding import semantic_content
from .teaching_plan import safe_fallback_plan, validate_teaching_plan, TeachingPlanValidationError

FIXTURES = json.loads((Path(__file__).parent / 'fixtures/material_intelligence.json').read_text(encoding='utf-8'))


class MaterialIntelligenceTests(SimpleTestCase):
    def model(self, key='finite_difference'):
        return build_understanding(key, FIXTURES[key])

    def test_structure_and_neighbors(self):
        model = self.model()
        self.assertEqual([page['number'] for page in model['pages']], [1, 2, 3])
        self.assertEqual(model['pages'][1]['previous'], 'page-1')
        self.assertEqual(model['pages'][1]['next'], 'page-3')
        bundle = grounding_bundle(model, 'finite difference', 2)
        self.assertEqual(len(bundle['pages']), 3)
        self.assertIn('5/1=5', bundle['excerpt'])

    def test_formula_and_table_provenance(self):
        model = self.model()
        formula = next(item for item in model['knowledge']['formulas'] if '[f(x+h)' in item['text'])
        self.assertEqual(formula['original'], formula['normalized'])
        self.assertEqual(formula['source_refs'][0]['number'], 2)
        self.assertEqual(formula['confidence'], 'uncertain')
        table = model['knowledge']['tables'][0]
        self.assertEqual(table['headers'], ['x', 'f(x)'])
        self.assertEqual(table['rows'], [['1', '3'], ['2', '8']])

    def test_finite_difference_has_real_worked_example(self):
        bundle = grounding_bundle(self.model(), 'Estimate finite difference', 2)
        representation, content, reason = semantic_content(bundle, 'WORKED_EXAMPLE')
        self.assertEqual(representation, 'WORKED_EXAMPLE')
        self.assertFalse(reason)
        self.assertIn('8-3=5', str(content))
        self.assertNotIn('Substitute the known information', str(content))
        item = SimpleNamespace(id='c', title='Finite difference', summary='', description='', difficulty='medium',
                               path=SimpleNamespace(subject='Mathematics', goal='Estimate a derivative'))
        plan = safe_fallback_plan(item, {'id': 'o', 'text': 'Calculate a finite difference.'}, bundle)
        self.assertIn('5/1=5', str(plan))

    def test_cross_subject_payloads(self):
        biology = self.model('biology')['knowledge']
        self.assertTrue(biology['concepts'])
        self.assertTrue(biology['entities'])
        self.assertEqual(biology['processes'][0]['steps'], ['Right heart', 'Lungs', 'Left heart', 'Body'])
        literature = self.model('literature')['knowledge']
        self.assertTrue(literature['quotations'])
        self.assertTrue(literature['comparisons'])
        self.assertTrue(literature['assessment'])
        architecture = self.model('computer_science')['knowledge']
        self.assertEqual(len(architecture['relationships']), 3)
        self.assertIn('PostgreSQL', architecture['relationships'][-1]['target'])

    def test_insufficient_representation_falls_back(self):
        representation, content, reason = semantic_content({'excerpt': 'Calculus finds derivatives.'}, 'WORKED_EXAMPLE')
        self.assertEqual(representation, 'GROUNDED_EXPLANATION')
        self.assertTrue(reason)
        self.assertNotIn('steps', content)

    def test_malformed_ai_rejected(self):
        for raw in ([], {'concepts': 'bad'}, {'unknown': []}, {'concepts': [{'text': 'invented', 'source_refs': [{'page_id': 'missing'}]}]}):
            with self.subTest(raw=raw), self.assertRaises(ValueError):
                validate_semantics(raw, self.model())

    def test_ai_cannot_invent_facts_in_valid_citation(self):
        item = {'text': 'Central differences are required.', 'quote': 'h is the spacing between neighboring x values.',
                'source_refs': [{'page_id': 'page-2'}], 'support': 'source'}
        with self.assertRaises(ValueError):
            validate_semantics({'concepts': [item]}, self.model())

    def test_ai_failure_preserves_deterministic_material(self):
        model = self.model()
        failed = understand_with_ai(model, Mock(side_effect=RuntimeError('offline')))
        self.assertEqual(failed['knowledge'], model['knowledge'])
        self.assertGreater(failed['ai_failed_windows'], 0)

    def test_valid_ai_classification_retains_canonical_provenance(self):
        quote = 'h is the spacing between neighboring x values.'
        chat = Mock(return_value=json.dumps({'variables': [{'text': quote, 'quote': quote,
                    'source_refs': [{'page_id': 'page-2', 'untrusted_extra': 'discard me'}], 'support': 'source'}]}))
        result = understand_with_ai(self.model(), chat)
        self.assertEqual(result['origin'], 'ai_assisted')
        classified = next(item for item in result['knowledge']['variables'] if item['confidence'] == 'ai_classified')
        self.assertEqual(classified['source_refs'], [{'page_id': 'page-2', 'number': 2}])
        self.assertEqual(chat.call_args.kwargs['task'], 'SOURCE_UNDERSTANDING')

    @patch('library.text_extractor._convert_pptx_to_pdf', return_value=None)
    @patch('library.text_extractor._render_slides_pillow', return_value=[])
    def test_pptx_boundaries_and_structured_table(self, *_):
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(3), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = 'x', 'f(x)'
        table.cell(1, 0).text, table.cell(1, 1).text = '1', '3'
        deck.slides.add_slide(deck.slide_layouts[6])
        data = io.BytesIO(); deck.save(data)
        parsed = extract_text_from_bytes(data.getvalue(), '.pptx')
        self.assertEqual(parsed['status'], 'success')
        self.assertEqual(len(parsed['pages']), 2)
        table_block = next(block for block in parsed['pages'][0]['blocks'] if block['kind'] == 'table')
        self.assertEqual(table_block['rows'], [['1', '3']])

    def test_pdf_pages_and_original_equation(self):
        import fitz
        document = fitz.open()
        document.new_page().insert_text((72, 72), 'y = (8 - 3) / 1')
        document.new_page().insert_text((72, 72), 'The estimated slope is 5.')
        parsed = extract_text_from_bytes(document.tobytes(), '.pdf')
        document.close()
        pages = parsed['pdf_data']['pages']
        self.assertEqual([page['number'] for page in pages], [1, 2])
        self.assertIn('y = (8 - 3) / 1', pages[0]['text'])

    def test_inference_and_enrichment_are_distinct(self):
        model = self.model()
        item = {'text': 'Prior slope knowledge may help.', 'quote': 'h is the spacing between neighboring x values.',
                'source_refs': [{'page_id': 'page-2'}], 'support': 'inferred'}
        semantic = validate_semantics({'prerequisites': [item]}, model)
        self.assertEqual(semantic['prerequisites'][0]['support'], 'inferred')
        model['knowledge']['prerequisites'] = [{**semantic['prerequisites'][0], 'support': 'enrichment'}]
        self.assertFalse(grounding_bundle(model, 'finite', 2)['knowledge']['prerequisites'])

    def test_teach_before_test_and_generic_steps(self):
        item = SimpleNamespace(id='c', title='Blood flow', summary='The heart pumps blood.', description='', difficulty='medium',
                               path=SimpleNamespace(subject='Biology', goal='Explain blood flow'))
        plan = safe_fallback_plan(item, {'id': 'o', 'text': 'Explain blood flow.'}, {'excerpt': item.summary})
        check = {'id': 'check', 'type': 'CHECK', 'representation': 'GROUNDED_EXPLANATION', 'interaction': 'SHORT_ANSWER',
                 'content': {'body': 'Answer the question.', 'prompt': 'What pumps blood?', 'expected_answer': 'heart', 'evidence_concepts': ['heart']}}
        plan['teaching_moments'].append(check)
        validate_teaching_plan(plan)
        plan['teaching_moments'].reverse()
        with self.assertRaises(TeachingPlanValidationError):
            validate_teaching_plan(plan)
        plan['teaching_moments'].reverse()
        plan['teaching_moments'][1]['content']['evidence_concepts'] = ['mitochondria']
        with self.assertRaises(TeachingPlanValidationError):
            validate_teaching_plan(plan)
        plan['teaching_moments'] = plan['teaching_moments'][:1]
        plan['teaching_moments'][0]['content']['steps'] = ['Substitute the known information.']
        with self.assertRaises(TeachingPlanValidationError):
            validate_teaching_plan(plan)

    def test_docx_tables_and_unknown_pagination(self):
        from docx import Document
        document = Document()
        document.add_heading('Measurements', 1)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = 'x', 'f(x)'
        table.cell(1, 0).text, table.cell(1, 1).text = '1', '3'
        output = io.BytesIO()
        document.save(output)
        parsed = extract_text_from_bytes(output.getvalue(), '.docx')
        self.assertEqual(parsed['status'], 'success')
        self.assertIsNone(parsed['pages'][0]['number'])
        self.assertEqual(parsed['pages'][0]['blocks'][-1]['rows'], [['1', '3']])

    def test_fingerprint_changes_with_material(self):
        self.assertEqual(self.model()['fingerprint'], self.model()['fingerprint'])
        self.assertNotEqual(self.model()['fingerprint'], build_understanding('finite_difference', 'changed')['fingerprint'])


class MaterialPersistenceTests(TestCase):
    @override_settings(JOURNEY_TEACHING_AI_ENABLED=False)
    def test_objective_to_player_and_submission_gate(self):
        from django.contrib.auth import get_user_model
        from library.models import Resource
        from learning.models import LearningPath, ConceptNode
        from learning.views import _get_teaching_session, _objective_activities, _session_data, submit_teaching_activity
        from learning.player import continue_player_stage
        user = get_user_model().objects.create_user(username='journey-material')
        resource = Resource.objects.create(owner=user, title='Finite difference')
        persist_understanding(resource, FIXTURES['finite_difference'])
        path = LearningPath.objects.create(user=user, title='Calculus', subject='Mathematics')
        concept = ConceptNode.objects.create(path=path, title='Finite difference', source_resource=resource, source_page=2)
        session = _get_teaching_session(concept, user)
        self.assertTrue(session.objectives[0]['source_refs'])
        activities = _objective_activities(session, user)
        check = next(item for item in activities if item['purpose'] == 'check')
        with self.assertRaisesMessage(ValueError, 'Complete the teaching'):
            submit_teaching_activity(concept, user, check['id'], {'text': 'guess'})
        _session_data(session)
        while not session.state['player'].get('active_activity_id'):
            self.assertTrue(continue_player_stage(session))
            _session_data(session)
        self.assertEqual(session.state['player']['active_activity_id'], check['id'])
        from learning.material_grounding import assessment_ready
        self.assertTrue(assessment_ready(session, check, activities))
        from learning.diagnostics import objective_trace
        with override_settings(DEBUG=True):
            trace = objective_trace(session, session.objectives[0]['id'])
        self.assertEqual(trace['teaching_plan_input']['source_fingerprint'], resource.source_understanding['fingerprint'])

    def test_cache_and_lazy_reuse(self):
        from django.contrib.auth import get_user_model
        from library.models import Resource
        resource = Resource.objects.create(owner=get_user_model().objects.create_user(username='material-test'), title='Finite difference')
        first = persist_understanding(resource, FIXTURES['finite_difference'])
        second = persist_understanding(resource, FIXTURES['finite_difference'], allow_ai=True)
        self.assertEqual(first, second)
        resource.refresh_from_db()
        self.assertEqual(resource.source_understanding, first)

    @override_settings(DEBUG=False)
    def test_diagnostics_disabled_in_production(self):
        from .diagnostics import objective_trace
        with self.assertRaises(PermissionError):
            objective_trace(None, 'o')
